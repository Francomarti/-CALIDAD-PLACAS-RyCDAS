#!/usr/bin/env python3
"""
Descarga el Excel de calidad/placas desde Google Drive, lo procesa y
regenera index.html. Pensado para correr desde GitHub Actions una vez
por semana (ver .github/workflows/update.yml), pero también se puede
correr a mano: python3 scripts/build.py
"""
import base64
import json
import os
import re
import sys
from datetime import date, datetime
from io import BytesIO

import openpyxl
import requests

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DRIVE_FILE_ID = "1MpWvhPVmLlbELGtCK0Nt7wO9yMGoZ4oQ"
TEMPLATE_PATH = os.path.join(REPO_ROOT, "template.html")
LOGO_PATH = os.path.join(REPO_ROOT, "assets", "logo.png")
OUTPUT_PATH = os.path.join(REPO_ROOT, "index.html")

CROP_NAMES = {"girasol", "maiz", "maíz", "soja", "sorgo", "pasturas", "garbanzos", "trigo"}

PASTURA_KEYWORDS = [
    "ALFALFA", "AGROPIRO", "FESTUCA", "RYE GRASS", "RAIGRAS", "TREBOL", "PASTURA",
    "MELILOTO", "CEBADILLA", "LOTUS", "VICIA", "TRITICALE", "PASTO OVILLO",
    "CEBADA FORRAJERA", "MOHA", "SUDAN", "RAYGRASS",
]

CROP_ORDER = ["Maiz", "Girasol", "Soja", "Sorgo", "Trigo", "Pasturas", "Garbanzos"]

# La hoja "maiz y gira carry" repite (con otro formato) filas que ya están en
# Girasol/Maiz, según confirmó Franco - la excluimos para no reprocesarla.
EXCLUDED_SHEETS = {"maiz y gira carry"}

# Color de fondo de la fila -> estado del lote (confirmado por Franco):
#   rosa/salmón = lote nuevo | gris = alerta (problema de PG u otro aviso) | blanco/sin color = carry
ESTADO_NUEVO_RGB = {"FFF7CAAC"}
ESTADO_ALERTA_RGB = {"FFAEABAB", "FF999999"}
ESTADO_CARRY_RGB = {"FFFFFFFF"}


def classify_estado(cell):
    fill = cell.fill
    if fill is None or fill.fill_type is None:
        return "carry"
    fg = fill.fgColor
    if fg is None or fg.type != "rgb":
        return None  # color de tema u otro caso no resuelto: lo dejamos sin clasificar
    rgb = fg.rgb
    if rgb in ESTADO_NUEVO_RGB:
        return "nuevo"
    if rgb in ESTADO_ALERTA_RGB:
        return "alerta"
    if rgb in ESTADO_CARRY_RGB:
        return "carry"
    return None
CROP_LABEL = {
    "Maiz": "Maíz", "Girasol": "Girasol", "Soja": "Soja", "Sorgo": "Sorgo",
    "Trigo": "Trigo", "Pasturas": "Pasturas/Forrajeras", "Garbanzos": "Garbanzos",
}

DEP_FIX = {
    "TANDIL CYO BASE AERE": "TANDIL CYO BASE AEREA",
    "TANDIL PROPIO BASE A": "TANDIL PROPIO BASE AEREA",
}


# ---------------------------------------------------------------- descarga
def download_xlsx(file_id: str) -> bytes:
    """Descarga un archivo de Google Drive compartido como 'cualquiera con el link'."""
    session = requests.Session()
    url = "https://drive.google.com/uc?export=download"
    resp = session.get(url, params={"id": file_id}, stream=True, timeout=60)

    # Archivos grandes: Drive interpone una página de advertencia con un token de confirmación.
    token = None
    for key, value in resp.cookies.items():
        if key.startswith("download_warning"):
            token = value
    if token is None and resp.headers.get("Content-Type", "").startswith("text/html"):
        m = re.search(r'confirm=([0-9A-Za-z_]+)', resp.text)
        if m:
            token = m.group(1)
    if token:
        resp = session.get(url, params={"id": file_id, "confirm": token}, stream=True, timeout=60)

    content = resp.content
    if not (content[:2] == b"PK"):  # los .xlsx son en realidad un .zip (firma PK)
        raise RuntimeError(
            f"La respuesta de Drive no parece un .xlsx válido (status={resp.status_code}, "
            f"content-type={resp.headers.get('Content-Type')}). "
            "Verificá que el archivo siga compartido como 'Cualquiera con el enlace puede ver'."
        )
    return content


# ---------------------------------------------------------------- helpers
def norm(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    return v


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(",", "."))
        except ValueError:
            return None
    return None


def to_date_iso(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat() if isinstance(v, date) and not isinstance(v, datetime) else v.date().isoformat()
    return None


def guess_crop(descripcion, sheet_name_norm):
    d = (descripcion or "").upper()
    if "MAIZ" in d or "MAÍZ" in d:
        return "Maiz"
    if "GIRASOL" in d:
        return "Girasol"
    if "SOJA" in d:
        return "Soja"
    if "SORGO" in d:
        return "Sorgo"
    if "TRIGO" in d:
        return "Trigo"
    if "GARBANZO" in d:
        return "Garbanzos"
    if any(k in d for k in PASTURA_KEYWORDS):
        return "Pasturas"
    # si la hoja se llama como un cultivo conocido, usamos eso como respaldo
    if sheet_name_norm in CROP_NAMES:
        return "Maiz" if sheet_name_norm == "maíz" else sheet_name_norm.capitalize()
    return "Otros"


def fix_dep(d):
    if not d:
        return d
    return DEP_FIX.get(d, d)


def sucursal_of(dep):
    if not dep:
        return "Sin dato"
    d = dep.upper()
    for name in ["AZUL", "BOLIVAR", "BOLÍVAR", "GBELGRANO", "BELGRANO", "OLAVARRIA", "OLAVARRÍA", "TANDIL"]:
        if name in d:
            if name in ("BOLIVAR", "BOLÍVAR"):
                return "Bolívar"
            if name in ("GBELGRANO", "BELGRANO"):
                return "Gral. Belgrano"
            if name in ("OLAVARRIA", "OLAVARRÍA"):
                return "Olavarría"
            return name.capitalize()
    return "Otra"


def parse_pg(pg_raw):
    if not pg_raw:
        return None, None
    s = str(pg_raw).strip()
    if s == "":
        return None, None
    m = re.match(r"^(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)$", s)
    if m:
        a, b = float(m.group(1)), float(m.group(2))
        return round((a + b) / 2, 1), s
    try:
        return float(s), s
    except ValueError:
        return None, s


def days_to(iso):
    if not iso:
        return None
    y, m, d = [int(x) for x in iso.split("-")]
    return (date(y, m, d) - date.today()).days


# ---------------------------------------------------------------- extracción
def read_workbook(content: bytes):
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    rows = []
    for sheet_name in wb.sheetnames:
        if sheet_name.strip().lower() in EXCLUDED_SHEETS:
            continue
        ws = wb[sheet_name]
        headers = {}
        for c in range(1, ws.max_column + 1):
            h = ws.cell(row=1, column=c).value
            if h is not None:
                headers[str(h).strip().lower()] = c

        def col(*names):
            for n in names:
                if n in headers:
                    return headers[n]
            return None

        c_art, c_desc = col("articulo"), col("descripcion")
        c_dep, c_lote = col("depósito", "deposito"), col("lote")
        c_venc, c_exist = col("vencimiento"), col("existencia")
        c_cond, c_pg = col("condicion"), col("pg")
        c_obs, c_placa = col("observaciones"), col("placas sugeridas", "placas sugerida")
        c_cold, c_energia = col("cold test"), col("test de energia")
        c_plant, c_pms = col("pantulas vigorosas"), col("pms")

        sheet_norm = sheet_name.strip().lower()
        for r in range(2, ws.max_row + 1):
            art = norm(ws.cell(row=r, column=c_art).value) if c_art else None
            desc = norm(ws.cell(row=r, column=c_desc).value) if c_desc else None
            if art is None and desc is None:
                continue
            row = dict(
                articulo=art, descripcion=desc,
                deposito=fix_dep(norm(ws.cell(row=r, column=c_dep).value) if c_dep else None),
                lote=norm(ws.cell(row=r, column=c_lote).value) if c_lote else None,
                vencimiento=to_date_iso(ws.cell(row=r, column=c_venc).value) if c_venc else None,
                existencia=to_num(ws.cell(row=r, column=c_exist).value) if c_exist else None,
                condicion=norm(ws.cell(row=r, column=c_cond).value) if c_cond else None,
                pg_raw=norm(ws.cell(row=r, column=c_pg).value) if c_pg else None,
                observaciones=norm(ws.cell(row=r, column=c_obs).value) if c_obs else None,
                placa=norm(ws.cell(row=r, column=c_placa).value) if c_placa else None,
                cold_test=norm(ws.cell(row=r, column=c_cold).value) if c_cold else None,
                test_energia=norm(ws.cell(row=r, column=c_energia).value) if c_energia else None,
                plantulas=norm(ws.cell(row=r, column=c_plant).value) if c_plant else None,
                pms=norm(ws.cell(row=r, column=c_pms).value) if c_pms else None,
                fuente=sheet_name,
                estado=classify_estado(ws.cell(row=r, column=1)),
            )
            row["cultivo"] = guess_crop(desc, sheet_norm)
            row["es_primaria"] = sheet_norm in CROP_NAMES
            rows.append(row)
    return rows


def dedup(rows):
    def completeness(row):
        keys = ["pg_raw", "observaciones", "placa", "cold_test", "test_energia", "plantulas", "pms",
                "vencimiento", "condicion"]
        return sum(1 for k in keys if row.get(k) not in (None, ""))

    best = {}
    for row in rows:
        key = (row["articulo"], row["lote"], row["deposito"])
        if key not in best:
            best[key] = row
            continue
        cur = best[key]
        cur_p = 0 if cur["es_primaria"] else 1
        new_p = 0 if row["es_primaria"] else 1
        if new_p < cur_p or (new_p == cur_p and completeness(row) > completeness(cur)):
            best[key] = row
    return list(best.values())


def aggregate(rows):
    for r in rows:
        r["sucursal"] = sucursal_of(r["deposito"])
        r["pg_val"], r["pg_disp"] = parse_pg(r["pg_raw"])
        r["dias_venc"] = days_to(r["vencimiento"])

    materials = {}
    for r in rows:
        key = (r["cultivo"], r["descripcion"], r["articulo"])
        materials.setdefault(key, []).append(r)

    by_crop = {}
    for (cultivo, desc, art), lots in materials.items():
        existencia_total = sum(l["existencia"] or 0 for l in lots)
        pg_vals = [l["pg_val"] for l in lots if l["pg_val"] is not None]
        placas = sorted(set(str(l["placa"]) for l in lots if l.get("placa")))
        cold_tests = sorted(set(str(l["cold_test"]) for l in lots if l.get("cold_test")))
        energias = sorted(set(str(l["test_energia"]) for l in lots if l.get("test_energia")))
        obs_set = sorted(set(str(l["observaciones"]) for l in lots if l.get("observaciones")))
        sucursales = sorted(set(l["sucursal"] for l in lots))
        min_dias = min([l["dias_venc"] for l in lots if l["dias_venc"] is not None], default=None)
        estados = sorted(set(l["estado"] for l in lots if l.get("estado")))

        lotes_trim = []
        for l in sorted(lots, key=lambda x: (x["sucursal"], x["deposito"] or "", str(x["lote"] or ""))):
            item = {
                "dep": l["deposito"], "suc": l["sucursal"], "lote": l["lote"],
                "venc": l["vencimiento"], "dias": l["dias_venc"], "ex": l["existencia"],
                "cond": l["condicion"], "pg": l["pg_disp"], "pgv": l["pg_val"],
                "obs": l["observaciones"], "placa": l["placa"], "estado": l.get("estado"),
            }
            if l.get("cold_test"):
                item["cold"] = l["cold_test"]
            if l.get("test_energia"):
                item["ener"] = l["test_energia"]
            if l.get("plantulas"):
                item["plant"] = l["plantulas"]
            if l.get("pms"):
                item["pms"] = l["pms"]
            lotes_trim.append(item)

        by_crop.setdefault(cultivo, []).append({
            "art": art, "desc": desc, "ex": existencia_total,
            "pgmin": min(pg_vals) if pg_vals else None,
            "pgmax": max(pg_vals) if pg_vals else None,
            "pgavg": round(sum(pg_vals) / len(pg_vals), 1) if pg_vals else None,
            "placas": placas, "cold": cold_tests, "ener": energias, "obs": obs_set,
            "suc": sucursales, "mindias": min_dias, "estados": estados, "lotes": lotes_trim,
        })

    for crop in by_crop:
        by_crop[crop].sort(key=lambda m: m["desc"] or "")

    ordered = {c: by_crop[c] for c in CROP_ORDER if c in by_crop}
    for c in by_crop:
        if c not in ordered:
            ordered[c] = by_crop[c]
    return ordered


def build():
    print(f"[{datetime.now().isoformat()}] Descargando planilla de Drive…")
    content = download_xlsx(DRIVE_FILE_ID)
    print(f"  {len(content)} bytes descargados")

    rows = read_workbook(content)
    print(f"  {len(rows)} filas crudas")
    rows = dedup(rows)
    print(f"  {len(rows)} filas tras dedup")

    by_crop = aggregate(rows)
    tot_mat = sum(len(v) for v in by_crop.values())
    tot_lot = sum(len(m["lotes"]) for v in by_crop.values() for m in v)
    print(f"  {tot_mat} materiales, {tot_lot} lotes")

    by_crop["__meta"] = {
        "fecha": date.today().strftime("%d/%m/%Y"),
        "origen": "Google Sheet (auto, semanal)",
        "totalMateriales": tot_mat,
        "totalLotes": tot_lot,
    }
    data_json = json.dumps(by_crop, ensure_ascii=False, separators=(",", ":"))

    with open(LOGO_PATH, "rb") as f:
        logo_b64 = base64.b64encode(f.read()).decode("ascii")

    with open(TEMPLATE_PATH, encoding="utf-8") as f:
        tpl = f.read()

    out = tpl.replace("__DATA__", data_json).replace("__LOGO_B64__", logo_b64)

    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        f.write(out)
    print(f"  index.html generado ({len(out)} bytes)")


if __name__ == "__main__":
    try:
        build()
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
